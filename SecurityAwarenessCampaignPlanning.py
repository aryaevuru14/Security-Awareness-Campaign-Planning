import docx

def create_campaign_plan_document():
    # Initialize the Word document
    doc = docx.Document()
    
    # Document Title
    doc.add_heading('Security Awareness Campaign Plan', level=0)
    
    # Section 1: Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        'This document outlines a comprehensive 4-week Security Awareness Campaign designed '
        'to educate employees on fundamental cybersecurity best practices, reduce human error, '
        'and foster a proactive security culture across the hypothetical organization.'
    )
    
    # Section 2: Campaign Objectives
    doc.add_heading('2. Campaign Objectives', level=1)
    p = doc.add_paragraph()
    p.add_run('• Reduce susceptibility to phishing attacks by 40%.\n')
    p.add_run('• Achieve 100% employee completion rate for core security training modules.\n')
    p.add_run('• Improve password hygiene and multi-factor authentication (MFA) adoption.\n')
    p.add_run('• Establish clear incident reporting channels for suspicious activities.')
    
    # Section 3: Target Audience & Scope
    doc.add_heading('3. Target Audience & Scope', level=1)
    doc.add_paragraph(
        'The campaign targets all internal employees, contractors, and remote workers. '
        'Content is tailored into general modules for all staff and advanced modules for technical teams.'
    )
    
    # Section 4: Key Topics
    doc.add_heading('4. Key Topics', level=1)
    doc.add_paragraph('1. Phishing & Social Engineering: Recognizing spear-phishing, smishing, and BEC (Business Email Compromise).')
    doc.add_paragraph('2. Password Hygiene: Creating strong passphrases and managing credentials securely.')
    doc.add_paragraph('3. Device & Remote Work Security: Screen locking, secure Wi-Fi usage, and physical security.')
    doc.add_paragraph('4. Incident Reporting: Knowing who to contact when a security anomaly is detected.')
    
    # Section 5: Communication Methods & Schedule
    doc.add_heading('5. Communication Methods & Schedule', level=1)
    doc.add_paragraph('• Week 1: Launch webinar and introductory leadership keynote.\n'
                      '• Week 2: Interactive phishing simulation and email newsletters.\n'
                      '• Week 3: Departmental workshops and interactive quizzes with rewards.\n'
                      '• Week 4: Campaign wrap-up, evaluation metrics review, and feedback survey.')
    
    # Section 6: Evaluation & Metrics
    doc.add_heading('6. Evaluation & Metrics', level=1)
    doc.add_paragraph(
        'Success will be measured using phishing simulation click rates, training completion percentages, '
        'and the volume of correctly reported security incidents.'
    )
    
    # Save the file
    doc.save('Security_Awareness_Campaign_Plan.docx')
    print("Document successfully created: Security_Awareness_Campaign_Plan.docx")

if __name__ == '__main__':
    create_campaign_plan_document()